import re
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.schemas import SlideOut

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
TITLE_COLOR = RGBColor(0x1A, 0x1A, 0x1A)
BODY_COLOR = RGBColor(0x33, 0x33, 0x33)


def sanitize_filename(title: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).strip() or "presentation"
    name = name.rstrip(". ")
    return f"{name}.pptx"


def build_pptx(title: str, slides: list[SlideOut]) -> bytes:
    prs = Presentation()
    prs.core_properties.title = title
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank = prs.slide_layouts[6]

    for slide in slides:
        pptx_slide = prs.slides.add_slide(blank)
        if slide.layout == "title":
            _render_title(pptx_slide, slide)
        elif slide.layout == "section":
            _render_section(pptx_slide, slide)
        elif slide.layout == "two_column":
            _render_two_column(pptx_slide, slide)
        else:
            _render_bullets(pptx_slide, slide)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _set_run(paragraph, text: str, size: Pt, bold: bool = False, color=TITLE_COLOR) -> None:
    run = paragraph.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _render_title(slide, data: SlideOut) -> None:
    box = _add_textbox(slide, Inches(1), Inches(2.4), Inches(11.333), Inches(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p, data.title, Pt(40), bold=True)

    subtitle = str(data.body.get("subtitle") or "")
    if subtitle:
        sub = _add_textbox(slide, Inches(1.5), Inches(4.1), Inches(10.333), Inches(1))
        sp = sub.text_frame.paragraphs[0]
        sp.alignment = PP_ALIGN.CENTER
        _set_run(sp, subtitle, Pt(22), color=BODY_COLOR)


def _render_section(slide, data: SlideOut) -> None:
    box = _add_textbox(slide, Inches(1), Inches(2.8), Inches(11.333), Inches(2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p, data.title, Pt(44), bold=True)


def _render_bullets(slide, data: SlideOut) -> None:
    title_box = _add_textbox(slide, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1))
    tp = title_box.text_frame.paragraphs[0]
    _set_run(tp, data.title, Pt(28), bold=True)

    bullets = data.body.get("bullets") or []
    if not isinstance(bullets, list):
        bullets = []
    body = _add_textbox(slide, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.text = str(item)
        p.font.size = Pt(20)
        p.font.color.rgb = BODY_COLOR


def _render_two_column(slide, data: SlideOut) -> None:
    title_box = _add_textbox(slide, Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
    _set_run(title_box.text_frame.paragraphs[0], data.title, Pt(28), bold=True)

    left_title = str(data.body.get("left_title") or "")
    right_title = str(data.body.get("right_title") or "")
    left_items = data.body.get("left") or []
    right_items = data.body.get("right") or []
    if not isinstance(left_items, list):
        left_items = []
    if not isinstance(right_items, list):
        right_items = []

    _column(slide, Inches(0.8), left_title, left_items)
    _column(slide, Inches(7.0), right_title, right_items)


def _column(slide, left, heading: str, items: list) -> None:
    head = _add_textbox(slide, left, Inches(1.5), Inches(5.5), Inches(0.6))
    if heading:
        _set_run(head.text_frame.paragraphs[0], heading, Pt(18), bold=True, color=BODY_COLOR)

    body = _add_textbox(slide, left, Inches(2.2), Inches(5.5), Inches(4.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(item)
        p.font.size = Pt(18)
        p.font.color.rgb = BODY_COLOR
